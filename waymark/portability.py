"""Bounded spreadsheet interchange. The application database remains canonical."""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
FIELDS = [
    "id", "version", "company", "role", "location", "employment_type", "level", "start_period",
    "watch_from", "watch_until", "notes", "reference_url", "source_url", "connector",
    "check_interval_hours", "monitoring_status", "historical_date", "historical_date_precision",
    "historical_date_meaning",
]
EVIDENCE_FIELDS = [
    "id", "target_id", "url", "excerpt", "retrieved_at", "kind", "date_value", "date_precision",
    "date_meaning", "match_status", "explanation",
]
ALIASES = {
    "employer": "company", "company_name": "company", "title": "role", "job_title": "role",
    "program": "role", "desired_role": "role", "region": "location", "country": "location",
    "cohort": "start_period", "start_date": "start_period", "intended_start": "start_period",
    "job_url": "reference_url", "prior_posting_url": "reference_url", "careers_url": "source_url",
    "watch_start": "watch_from", "watch_end": "watch_until",
}
_DANGEROUS = re.compile(r"^[\s]*[=+@-]|^[\t\r\n]")


class ImportError(ValueError):
    """A safe, user-facing import validation error."""


def _header(value: Any) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return ALIASES.get(text, text)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _read(filename: str, data: bytes) -> tuple[list[str], list[list[Any]]]:
    if len(data) > MAX_BYTES:
        raise ImportError("File is too large. Maximum upload is 5 MB.")
    if filename.lower().endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ImportError("Save the CSV as UTF-8 and try again.") from exc
        try:
            if text.strip():
                # Sniffer may mistake an apostrophe used for formula
                # neutralization for the CSV quote character. Detect only the
                # delimiter and keep RFC 4180's double quote contract.
                sniffed = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
                dialect = csv.excel()
                dialect.delimiter = sniffed.delimiter
            else:
                dialect = csv.excel()
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        rows = []
        try:
            for row in reader:
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ImportError(f"Import at most {MAX_ROWS:,} targets at a time.")
        except csv.Error as exc:
            raise ImportError("Could not parse the CSV.") from exc
    elif filename.lower().endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                    raise ImportError("The expanded workbook exceeds the safe import size.")
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book["Targets"] if "Targets" in book.sheetnames else book.active
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ImportError(f"Import at most {MAX_ROWS:,} targets at a time.")
                if sheet.max_column and sheet.max_column > 200:
                    raise ImportError("Import at most 200 columns at a time.")
                rows = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ImportError("Workbooks with formulas are not imported. Paste values first.")
                    rows.append([cell.value for cell in row])
                    if len(rows) > MAX_ROWS + 1:
                        raise ImportError(f"Import at most {MAX_ROWS:,} targets at a time.")
            finally:
                book.close()
        except ImportError:
            raise
        except Exception as exc:
            raise ImportError("Could not read this XLSX workbook.") from exc
    else:
        raise ImportError("Choose a CSV or XLSX file.")
    if not rows:
        raise ImportError("The file is empty.")
    columns = [_text(value).strip() for value in rows[0]]
    if len(columns) > 200:
        raise ImportError("Import at most 200 columns at a time.")
    if len(set(columns)) != len(columns):
        raise ImportError("Column names must be unique.")
    return columns, rows[1:]


def preview_import(filename: str, data: bytes, mapping: dict | None = None) -> dict:
    columns, raw_rows = _read(filename, data)
    source_indexes = {column: index for index, column in enumerate(columns)}
    if mapping:
        invalid = [field for field, column in mapping.items()
                   if field not in FIELDS or not isinstance(column, str) or column not in source_indexes]
        if invalid:
            raise ImportError("Column mapping contains an unknown field or source column.")
        indexes = {field: source_indexes[column] for field, column in mapping.items()}
    else:
        indexes = {_header(column): index for index, column in enumerate(columns) if _header(column) in FIELDS}
    rows, errors, ids = [], [], set()
    for number, raw in enumerate(raw_rows, start=2):
        if not any(value is not None and str(value).strip() for value in raw):
            continue
        row = {field: raw[index] if index < len(raw) else None for field, index in indexes.items()}
        escape_index = source_indexes.get("__escaped_fields")
        if escape_index is not None and escape_index < len(raw):
            try:
                escaped = json.loads(raw[escape_index] or "[]")
                for field in escaped if isinstance(escaped, list) else []:
                    value = row.get(field)
                    if isinstance(value, str) and value.startswith("'") and _DANGEROUS.search(value[1:]):
                        row[field] = value[1:]
            except (ValueError, TypeError):
                errors.append({"row": number, "message": "Invalid export escape metadata."})
                continue
        try:
            for field in tuple(row):
                value = row[field]
                if field in ("version", "check_interval_hours"):
                    if value in (None, ""):
                        del row[field]
                    else:
                        row[field] = int(value) if field == "version" else float(value)
                elif field in ("watch_from", "watch_until"):
                    row[field] = _text(value)[:10] or None
                    if row[field]:
                        date.fromisoformat(row[field])
                elif field == "historical_date":
                    row[field] = _text(value) or None
                else:
                    row[field] = _text(value)
                    if row[field] == "" and field not in ("notes", "location", "start_period", "reference_url", "source_url"):
                        del row[field]
            if not str(row.get("company", "")).strip() or not str(row.get("role", "")).strip():
                raise ValueError("Company and role are required. Map their columns before importing.")
            if any(len(str(value or "")) > 10000 for value in row.values()):
                raise ValueError("A cell exceeds the 10,000 character limit.")
            if row.get("watch_from") and row.get("watch_until") and row["watch_from"] > row["watch_until"]:
                raise ValueError("Watch start must not be after watch end.")
            if row.get("id"):
                if row["id"] in ids:
                    raise ValueError("Duplicate target ID in this import.")
                ids.add(row["id"])
            # Importing a file never enables monitoring implicitly.
            row["monitoring_status"] = "draft"
            rows.append(row)
        except (ValueError, TypeError) as exc:
            errors.append({"row": number, "message": str(exc)})
    return {"columns": columns, "rows": rows, "errors": errors, "existing_ids": sorted(ids)}


def _safe_csv(value: Any) -> tuple[str, bool]:
    text = _text(value)
    escaped = bool(_DANGEROUS.search(text))
    return ("'" + text if escaped else text), escaped


def _sheet(book: Workbook, name: str, fields: list[str], rows: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(name)
    sheet.append(fields)
    for record in rows:
        sheet.append([_text(record.get(field)) for field in fields])
        for cell in sheet[sheet.max_row]:
            # Force text, including values beginning '=', so Excel never executes input.
            cell.data_type = "s"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="143D36")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, field in enumerate(fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = 45 if field in ("url", "source_url", "excerpt", "notes") else 24
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_targets(targets: list[dict], evidence: list[dict], format: str) -> tuple[bytes, str, str]:
    if format == "csv":
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(buffer, fieldnames=FIELDS + ["__escaped_fields"])
        writer.writeheader()
        for target in targets:
            row, escaped_fields = {}, []
            for field in FIELDS:
                row[field], escaped = _safe_csv(target.get(field))
                if escaped:
                    escaped_fields.append(field)
            row["__escaped_fields"] = json.dumps(escaped_fields, separators=(",", ":"))
            writer.writerow(row)
        return buffer.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "waymark-targets.csv"
    if format == "xlsx":
        book = Workbook()
        book.remove(book.active)
        _sheet(book, "Targets", FIELDS, targets)
        _sheet(book, "Evidence", EVIDENCE_FIELDS, evidence)
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "waymark-targets.xlsx"
    raise ValueError("Export format must be csv or xlsx.")
