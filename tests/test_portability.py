import csv
import io

from openpyxl import Workbook, load_workbook
import pytest

from waymark.portability import export_targets, preview_import


@pytest.mark.parametrize("value", ['=HYPERLINK("https://example.com")', '+123', '-123', '@SUM(1)', '\t=1'])
def test_csv_formula_escaping_round_trip(value):
    payload, _, _ = export_targets([{"company": "Example", "role": "Engineer", "notes": value}], [], "csv")
    record = next(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    assert record["notes"] == "'" + value
    preview = preview_import("targets.csv", payload)
    assert preview["errors"] == []
    assert preview["rows"][0]["notes"] == value
    assert preview["rows"][0]["monitoring_status"] == "draft"


def test_xlsx_exports_formula_like_values_as_text_and_rejects_formulas():
    payload, _, _ = export_targets([{"company": "Example", "role": "Engineer", "notes": "=1+1"}], [], "xlsx")
    book = load_workbook(io.BytesIO(payload))
    sheet = book["Targets"]
    notes_column = [cell.value for cell in sheet[1]].index("notes") + 1
    assert sheet.cell(2, notes_column).data_type == "s"
    book.close()
    assert preview_import("targets.xlsx", payload)["rows"][0]["notes"] == "=1+1"
    malicious = Workbook()
    malicious.active.append(["company", "role"])
    malicious.active.append(["=1+1", "Engineer"])
    buffer = io.BytesIO()
    malicious.save(buffer)
    malicious.close()
    with pytest.raises(ValueError, match="formulas"):
        preview_import("targets.xlsx", buffer.getvalue())


def test_invalid_mapping_is_a_validation_error():
    with pytest.raises(ValueError, match="mapping"):
        preview_import("targets.csv", b"company,role\nExample,Engineer\n", {"company": []})
